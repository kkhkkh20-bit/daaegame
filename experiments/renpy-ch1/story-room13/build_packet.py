"""Export the story packet for reading. No game launch, lint, or QA."""
import html
import json
import re
import shutil
import zipfile
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from authoring_data import make_data

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'exports'
OUT.mkdir(parents=True,exist_ok=True)
data=make_data()
(ROOT/'narrative.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

scene_source=(ROOT/'SCENES.txt').read_text(encoding='utf-8')
scenes=[]
for part in re.split(r'(?m)^@@ ',scene_source)[1:]:
    heading,_,body=part.partition('\n')
    fields=[x.strip() for x in heading.split('|')]
    scenes.append({'id':fields[0],'heading':heading,'body':body.strip()})
data['scenes']=scenes
(ROOT/'narrative.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

wb=Workbook()
wb.remove(wb.active)
def sheet(title,headers,rows,widths):
    ws=wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append([json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v for v in row])
    ws.freeze_panes='A2'
    ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:
        c.fill=PatternFill('solid',fgColor='163A3B')
        c.font=Font(name='맑은 고딕',color='FFFFFF',bold=True)
        c.alignment=Alignment(vertical='center',wrap_text=True)
    ws.row_dimensions[1].height=32
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font=Font(name='맑은 고딕',size=10)
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            if cell.row%2==0:cell.fill=PatternFill('solid',fgColor='F1F5F3')
        ws.row_dimensions[row[0].row].height=66 if title!='장면 대사' else 180
    for i,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=width
    ws.sheet_view.zoomScale=85
    ws.print_options.horizontalCentered=True
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape'
    ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1
    ws.page_setup.fitToHeight=0
    ws.print_title_rows='1:1'
    return ws

sheet('읽기 안내',['항목','내용'],[
 ['상태','전체 본편 집필안. 현재 게임에 연결되지 않았음. 플레이/저장 검사는 실행하지 않음.'],
 ['범인','토끼 할머니. 엄마와 이미 적대하며 편지로 아빠를 유인. 촌장은 희생 집행의 책임자.'],
 ['피해자','부엉이 할아버지. 초롱은 할머니의 아들인 원작 인연을 유지해 엄마 연락책으로 남음.'],
 ['기본 일정','기본 주간 위치이며 사건/구금/실종/대피 상태가 우선함. 미래 일정 전체를 게임 UI에 노출하지 않음.'],
 ['행동 비용','질문/조사/자료 읽기/제시는 시간 소모 없음. 실제 동행/구출/수리/관측만 큰 시간대를 넘김.'],
 ['주민 수','초창기 명단 35명 + 가족 3명/너울/세련/돌목의 별도 역할. 모두 새 그림을 제작했다는 뜻은 아님.'],
 ['대사 범위',f'주요 장면 {len(scenes)}개. 35명의 모든 일상 대사를 최종 녹음본으로 작성한 상태는 아님.'],
 ['규칙','회차 수 잠금 없음. 현재 회차 증거만 제시. 사건 예방/할머니 조기 구금이 실제 줄거리를 바꿈.']], [22,110])

sheet('14일 본편',['일차','요일','장 제목','진행','중심 질문','장면'],
 [(x['day'],x['weekday'],x['title'],x['objective'],x['question'],' '.join(x['scenes'])) for x in data['day_flow']], [9,8,24,48,48,36])
truth_text=(ROOT/'CASE13_TRUTH.txt').read_text(encoding='utf-8')
truth_rows=re.findall(r'(?m)^(\d\d:\d\d\S*) (.+)$',truth_text)
sheet('4일 진실 시간표',['작가용 시각','실제 행동·목격·남는 흔적'],truth_rows,[23,135])
sheet('주민 35명',['ID','이름','동물','일상 목적','지키려는 것','실제로 아는 것','협조 조건','거절 조건','사건 후 변화','실행본 연결 ID'],
 [(x['id'],x['name'],x['species'],x['work'],x['protects'],x['private_knowledge'],x['opens_with'],x['refuses_when'],x['post_case_change'],x['runtime_id']) for x in data['cast']], [20,22,15,32,34,44,46,40,52,25])
week_rows=[]
daily_rows=[]
for person in data['cast']:
    for dow,triple in enumerate(person['weekly_locations']):
        week_rows.append([person['name'],'월화수목금토일'[dow],*[data['places'][p] for p in triple],person['work']])
    for day in range(1,15):
        triple=person['weekly_locations'][(day-1)%7]
        for period,p in enumerate(triple):
            note='상태별 분기표가 우선'
            shown=data['places'][p]
            if person['id']=='owlgp' and day>=5:
                shown='기본 경로: 사망 / 예방 경로: 안전 거처'
                note='예방하지 않은 경로에서 기본 주간 출석 금지'
            if person['id']=='grandma' and day>=8:
                shown='자유: 여관 / 입증: 구금시설'
                note='구금이면 여관/제단 자유 이동 및 명령 불가'
            if day==4 and period==2 and person['id'] in ['grandma','owlgp','bami','nabi','tok','tori','karo']:
                note='CASE13_TRUTH의 분 단위 동선과 S13/S43 조건 우선'
            if day>=8 and person['id'] in ['teacher','nabi','bori','wanggu']:
                shown=data['places']['school'];note='재난 이후 대피/식사/운반. 자원·구금 상태 확인'
            daily_rows.append([day,'월화수목금토일'[(day-1)%7],['오전','오후','저녁'][period],person['name'],shown,person['work'],note])
sheet('기본 주간 일정',['이름','요일','오전','오후','저녁','그곳에서 하는 일'],week_rows,[22,10,25,25,25,50])
sheet('14일 등장 기준',['일차','요일','시간대','이름','기본 장소 또는 상태','행동','우선 적용 조건'],daily_rows,[10,9,13,22,43,40,60])
sheet('상태별 일정 변경',['규칙','우선도','조건','인물','위치','행동'],
 [(x['id'],x['priority'],x['when'],x['actors'],x['location'],x['behavior']) for x in data['schedule_overrides']], [28,10,55,35,20,70])
sheet('증거와 한계',['자료 ID','이름','유형','확인할 수 있는 것','확인하지 못하는 것','장면'],
 [(x['id'],x['name'],x['kind'],x['establishes'],x['does_not_establish'],' '.join(x['scenes'])) for x in data['evidence']], [12,35,17,58,58,26])
sheet('조사 활동',['ID','활동','시작일','종료일','장소','시간대 진행','선행 조건','완료 결과','장면'],
 [(x['id'],x['title'],x['start_day'],x['end_day'],x['place'],x['time_cost'],x['requires'],x['grants_after_completion'],x['scene']) for x in data['actions']], [12,42,10,10,20,14,50,45,12])
sheet('주민에게 제시',['ID','상대','관련 자료','협조 조건','열리는 주제','주제 진행 후 결과','위험과 한계'],
 [(x['id'],x['recipient'],x['items'],x['requires'],x['topic'],x['reward_after_topic'],x['risk']) for x in data['presentations']], [22,20,24,35,35,35,70])
sheet('대결 논점',['ID','주장','전체 논증 자료 묶음','대안과 한계','반격','해결 결과','장면'],
 [(x['id'],x['assertion'],x['support_sets'],x['alternative'],x['counter'],x['result'],x['scene']) for x in data['claims']], [10,45,60,70,48,32,12])
sheet('엔딩 분기',['ID','이름','현재 회차 조건','남기는 지식','장면'],
 [(x['id'],x['title'],x['when'],x['knowledge'],x['scene']) for x in data['endings']], [10,28,100,60,12])
sheet('장면 대사',['ID','장면','대본과 선택'],[(s['id'],s['heading'],s['body']) for s in scenes],[10,58,130])
sheet('상태 근거',['상태','현재 회차에서 얻는 경로'],list(data['fact_sources'].items()),[35,115])
wb.save(OUT/'다람탐정_13호_스토리와분기표.xlsx')

docs=[('MAIN_STORY.txt','전체 본편'),('CASE13_TRUTH.txt','밀실 사건의 진실'),('SCENES.txt','장면 대본'),('IMPLEMENTATION.txt','게임 연결 설계')]
nav=''.join(f'<a href="#doc{i}">{title}</a>' for i,(_,title) in enumerate(docs))
blocks=''.join(f'<section id="doc{i}"><h2>{title}</h2><pre>{html.escape((ROOT/name).read_text(encoding="utf-8"))}</pre></section>' for i,(name,title) in enumerate(docs))
diagram='''<svg viewBox="0 0 920 420" role="img" aria-label="13호 정문과 서비스 통로, 식당 알리바이 동선의 개념도">
<defs><marker id="arrow" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto"><path d="M0 0L8 4L0 8Z" fill="#ae6241"/></marker></defs>
<g fill="#f4eee0" stroke="#547d75" stroke-width="2"><rect x="20" y="30" width="175" height="130" rx="12"/><rect x="255" y="30" width="290" height="185" rx="12"/><rect x="605" y="30" width="175" height="90" rx="12"/><rect x="605" y="180" width="175" height="85" rx="12"/><rect x="255" y="305" width="290" height="85" rx="12"/><rect x="20" y="280" width="175" height="110" rx="12"/></g>
<g font-family="sans-serif" font-size="19" fill="#203d39" text-anchor="middle"><text x="107" y="73">2층 복도</text><text x="107" y="104" font-size="15">정문만 감시</text><text x="400" y="68">13호실</text><text x="400" y="100" font-size="15">안쪽 빗장 / 피해자</text><text x="470" y="170" font-size="15">침대 아래: 다람</text><text x="692" y="70">옛 은신처</text><text x="692" y="99" font-size="14">엄마의 과거 생활 흔적</text><text x="692" y="218">창고 · 계단</text><text x="692" y="247" font-size="14">이동 약 45초 이상</text><text x="400" y="342">주방 · 세탁실</text><text x="400" y="371" font-size="14">연속 목격이 끊기는 구간</text><text x="107" y="320">1층 식당</text><text x="107" y="350" font-size="14">첫 비명 때</text><text x="107" y="373" font-size="14">할머니의 참인 알리바이</text><text x="825" y="165" font-size="14">처마의 밤이</text><text x="825" y="188" font-size="12">침대 아래만 관찰</text></g>
<g fill="none" stroke="#ae6241" stroke-width="3" marker-end="url(#arrow)"><path d="M195 335H255"/><path d="M545 345H570V220H605"/><path d="M692 180V120"/><path d="M605 75H545"/></g><path d="M195 125H255" stroke="#547d75" stroke-width="4"/><text x="221" y="148" font-size="12" text-anchor="middle">잠긴 정문</text><path d="M780 165H545" stroke="#547d75" stroke-width="2" stroke-dasharray="5 6"/><text x="30" y="240" font-size="14" fill="#687b75">개념도: 실제 배경은 문서의 규격·시야를 따라 제작</text></svg>'''
page=f'''<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>다람탐정 · 13호 본편 시나리오</title>
<style>body{{margin:0;background:#f7f4ec;color:#203d39;font-family:system-ui,"Noto Sans KR",sans-serif;line-height:1.85}}header{{background:#173e3c;color:#f4efdf;padding:48px max(24px,calc((100vw - 1060px)/2))}}h1{{font-size:clamp(28px,4vw,46px);line-height:1.4}}.tag{{color:#f0c898}}nav{{display:flex;gap:14px;flex-wrap:wrap;position:sticky;top:0;background:#e6ece6;padding:12px 24px;border-bottom:1px solid #a6b9af}}nav a{{color:#214c46;text-decoration:none;font-weight:650}}main{{max-width:1060px;margin:auto;padding:28px 24px}}section{{scroll-margin-top:75px;padding-top:28px;border-bottom:1px solid #bbc9bf}}pre{{white-space:pre-wrap;overflow-wrap:anywhere;font:inherit;line-height:1.95}}svg{{width:100%;height:auto;background:#fffdf6;border:1px solid #b8c7bc;border-radius:12px}}.note{{padding:18px 22px;background:#e5ece4;border-left:4px solid #688c75}}h2{{font-size:27px}}@media print{{nav{{display:none}}header{{padding:20px;color:#173e3c;background:white}}section{{break-before:page}}body{{font-size:10pt}}}}</style>
<header><div class="tag">전체 스포일러 · 집필본 v1 · 2026.10.12</div><h1>눈 아래의<br>열세 번째 방</h1><p>다람탐정 1편 · 14일 본편 / 살인 누명 / 가족 구출 / 재난 속 대결</p></header><nav>{nav}</nav><main>
<p class="note">스토리와 주요 대사·분기 설계를 작성한 문서입니다. 현재 게임 실행본에는 연결하지 않았습니다. 기존 공개 버전과 저장은 변경하지 않았고, 플레이 검사는 실행하지 않았습니다.</p>
<p>주요 장면 {len(scenes)}개 · 원작 주민 35명 · 자료 {len(data['evidence'])}개 · 조사 활동 {len(data['actions'])}개 · 엔딩 7종</p>{diagram}{blocks}</main></html>'''
(OUT/'다람탐정_13호_전체시나리오.html').write_text(page,encoding='utf-8')
# Document printing only: no game page, gameplay driver, or QA is opened.
from playwright.sync_api import sync_playwright
with sync_playwright() as p:
    browser=p.chromium.launch(executable_path='/usr/bin/chromium',headless=True,args=['--no-sandbox'])
    tab=browser.new_page()
    tab.set_content(page,wait_until='load')
    tab.evaluate('document.fonts.ready')
    tab.pdf(path=str(OUT/'다람탐정_13호_전체시나리오.pdf'),format='A4',print_background=True,
            margin={'top':'15mm','bottom':'17mm','left':'12mm','right':'12mm'},
            display_header_footer=True,header_template='<span></span>',
            footer_template='<div style="font:9px sans-serif;width:100%;text-align:center;color:#777">ROOM 13 · STORY DRAFT v1 — <span class="pageNumber"></span> / <span class="totalPages"></span></div>')
    browser.close()
for name in ['README.txt','MAIN_STORY.txt','CASE13_TRUTH.txt','SCENES.txt','IMPLEMENTATION.txt','narrative.json']:
    shutil.copy2(ROOT/name,OUT/name)
zip_path=OUT.parent/'다람탐정_13호_전체스토리.zip'
with zipfile.ZipFile(zip_path,'w',zipfile.ZIP_DEFLATED) as z:
    for path in sorted(OUT.iterdir()):z.write(path,path.name)
manifest={'status':'authored_not_connected_to_runtime','legacy_residents':len(data['cast']),
          'main_scenes':len(scenes),'evidence':len(data['evidence']),'actions':len(data['actions']),
          'endings':len(data['endings']),'artifacts':[str(p) for p in sorted(OUT.iterdir())],
          'zip':str(zip_path),'gameplay_checks':'not_run_at_user_request'}
(ROOT/'PACKET_STATUS.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({k:v for k,v in manifest.items() if k!='artifacts'},ensure_ascii=False))
